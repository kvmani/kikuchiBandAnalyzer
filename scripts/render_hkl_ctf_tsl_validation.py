"""Render a one-slide validation comparison for HKL CTF to TSL outputs."""

from __future__ import annotations

import argparse
import logging
from pathlib import Path
import sys

import h5py
import matplotlib
import numpy as np
from orix.crystal_map import Phase
from orix.plot import IPFColorKeyTSL
from orix.quaternion import Orientation, Rotation
from orix.vector import Vector3d
from PIL import Image
from pptx import Presentation

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from kikuchiBandAnalyzer.io.hkl_ctf_to_tsl import parse_ctf_file  # noqa: E402


def build_parser() -> argparse.ArgumentParser:
    """Build a command-line parser for the validation renderer.

    Returns:
        Configured argument parser.
    """

    parser = argparse.ArgumentParser(
        description="Render CTF/ANG/H5 IPF-X comparison against the HKL PPT reference."
    )
    parser.add_argument("--ctf", type=Path, default=Path("testData/hkl_ctf_test_data/Subset.ctf"))
    parser.add_argument(
        "--ang",
        type=Path,
        default=Path("testData/hkl_ctf_test_data/converted_tsl/Subset_HKL_TSL.ang"),
    )
    parser.add_argument(
        "--h5",
        type=Path,
        default=Path("testData/hkl_ctf_test_data/converted_tsl/Subset_HKL_TSL.h5"),
    )
    parser.add_argument(
        "--reference-pptx",
        type=Path,
        default=Path("testData/hkl_ctf_test_data/AcquisitionDetails.pptx"),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("testData/hkl_ctf_test_data/validation"),
    )
    return parser


def main() -> None:
    """Render validation map panels and a comparison PNG.

    Returns:
        None.
    """

    logging.basicConfig(level=logging.INFO, format="%(levelname)s:%(name)s:%(message)s")
    args = build_parser().parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    reference = extract_reference_ipf_x(args.reference_pptx, args.output_dir / "hkl_reference_ipf_x.png")
    ctf_rgb = ctf_ipf_x(args.ctf)
    ang_rgb = ang_ipf_x(args.ang)
    h5_rgb = h5_ipf_x(args.h5)
    write_image(ctf_rgb, args.output_dir / "ctf_ipf_x.png")
    write_image(ang_rgb, args.output_dir / "ang_ipf_x.png")
    write_image(h5_rgb, args.output_dir / "h5_ipf_x.png")
    render_comparison_png(
        reference_path=reference,
        ctf_rgb=ctf_rgb,
        ang_rgb=ang_rgb,
        h5_rgb=h5_rgb,
        output_path=args.output_dir / "hkl_ctf_tsl_ipf_x_comparison.png",
    )


def ctf_ipf_x(ctf_path: Path) -> np.ndarray:
    """Build an IPF-X RGB map directly from CTF Euler angles.

    Parameters:
        ctf_path: Source CTF path.

    Returns:
        RGB image array shaped ``(ny, nx, 3)``.
    """

    ctf = parse_ctf_file(ctf_path)
    eulers = np.column_stack(
        [ctf.table["Euler1"], ctf.table["Euler2"], ctf.table["Euler3"]]
    )
    valid = ctf.table["Phase"] > 0
    return eulers_to_ipf_x(np.deg2rad(eulers), valid=valid, shape=(ctf.ny, ctf.nx))


def ang_ipf_x(ang_path: Path) -> np.ndarray:
    """Build an IPF-X RGB map from generated ANG Euler angles.

    Parameters:
        ang_path: Generated ANG path.

    Returns:
        RGB image array shaped ``(ny, nx, 3)``.
    """

    lines = ang_path.read_text(encoding="utf-8").splitlines()
    nrows = 0
    ncols = 0
    header_end = 0
    for index, line in enumerate(lines):
        if line.startswith("# NROWS:"):
            nrows = int(line.split(":", 1)[1])
        elif line.startswith("# NCOLS_EVEN:"):
            ncols = int(line.split(":", 1)[1])
        elif line.startswith("# HEADER: End"):
            header_end = index + 1
            break
    rows = [line.split() for line in lines[header_end:] if len(line.split()) >= 8]
    data = np.asarray([[float(value) for value in row[:8]] for row in rows], dtype=np.float64)
    valid = data[:, 7] >= 0
    return eulers_to_ipf_x(data[:, :3], valid=valid, shape=(nrows, ncols))


def h5_ipf_x(h5_path: Path) -> np.ndarray:
    """Build an IPF-X RGB map from generated HDF5 Euler datasets.

    Parameters:
        h5_path: Generated HDF5 path.

    Returns:
        RGB image array shaped ``(ny, nx, 3)``.
    """

    with h5py.File(h5_path, "r") as handle:
        scan = next(key for key in handle.keys() if key not in {"Manufacturer", "Version"})
        header = handle[f"{scan}/EBSD/Header"]
        data = handle[f"{scan}/EBSD/Data"]
        nx = int(header["nColumns"][()][0])
        ny = int(header["nRows"][()][0])
        eulers = np.column_stack(
            [data["Phi1"][()], data["Phi"][()], data["Phi2"][()]]
        )
        valid = data["Phase"][()] >= 0
    try:
        import kikuchipy as kp

        signal = kp.load(h5_path, lazy=True)
        logging.info("kikuchipy loaded %s with axes %s", h5_path, signal.axes_manager)
    except Exception as exc:  # pragma: no cover - validation still works without optional load.
        logging.warning("kikuchipy load validation failed for %s: %s", h5_path, exc)
    return eulers_to_ipf_x(eulers, valid=valid, shape=(ny, nx))


def eulers_to_ipf_x(eulers_rad: np.ndarray, valid: np.ndarray, shape: tuple[int, int]) -> np.ndarray:
    """Convert Bunge Euler angles to an IPF-X RGB map using orix TSL coloring.

    Parameters:
        eulers_rad: Euler angles in radians shaped ``(n_pixels, 3)``.
        valid: Boolean mask indicating indexed pixels.
        shape: Output map shape as ``(ny, nx)``.

    Returns:
        RGB image array shaped ``(ny, nx, 3)``.
    """

    phase = Phase(name="Cr", space_group=225)
    rotations = Rotation.from_euler(eulers_rad, direction="lab2crystal", degrees=False)
    orientations = Orientation(rotations.data, symmetry=phase.point_group)
    key = IPFColorKeyTSL(phase.point_group, direction=Vector3d.xvector())
    rgb = key.orientation2color(orientations)
    rgb = np.asarray(rgb, dtype=np.float32)
    rgb[~valid] = np.array([0.18, 0.18, 0.18], dtype=np.float32)
    return rgb.reshape(shape[0], shape[1], 3)


def extract_reference_ipf_x(pptx_path: Path, output_path: Path) -> Path:
    """Extract the first IPF-X image from the acquisition details deck.

    Parameters:
        pptx_path: Acquisition details PPTX path.
        output_path: Destination PNG path.

    Returns:
        Path to the extracted reference image.
    """

    presentation = Presentation(pptx_path)
    for shape in presentation.slides[1].shapes:
        if getattr(shape, "shape_type", None) == 13:
            output_path.write_bytes(shape.image.blob)
            return output_path
    raise ValueError(f"No picture found on slide 2 of {pptx_path}.")


def write_image(rgb: np.ndarray, output_path: Path) -> None:
    """Write a float RGB image to disk.

    Parameters:
        rgb: Float RGB image in the range ``[0, 1]``.
        output_path: Destination PNG path.

    Returns:
        None.
    """

    payload = np.clip(rgb * 255.0, 0, 255).astype(np.uint8)
    Image.fromarray(payload).resize((300, 420), Image.Resampling.NEAREST).save(output_path)


def render_comparison_png(
    reference_path: Path,
    ctf_rgb: np.ndarray,
    ang_rgb: np.ndarray,
    h5_rgb: np.ndarray,
    output_path: Path,
) -> None:
    """Render the four-panel comparison PNG.

    Parameters:
        reference_path: Extracted HKL reference IPF-X image path.
        ctf_rgb: CTF-derived RGB map.
        ang_rgb: ANG-derived RGB map.
        h5_rgb: HDF5-derived RGB map.
        output_path: Destination PNG path.

    Returns:
        None.
    """

    fig, axes = plt.subplots(1, 4, figsize=(14, 4.8), constrained_layout=True)
    panels = [
        ("HKL PPT reference IPF-X", np.asarray(Image.open(reference_path).convert("RGB"))),
        ("Parsed CTF IPF-X", ctf_rgb),
        ("Generated ANG IPF-X", ang_rgb),
        ("Generated H5/OH5 IPF-X", h5_rgb),
    ]
    for axis, (title, image) in zip(axes, panels):
        axis.imshow(image, interpolation="nearest")
        axis.set_title(title, fontsize=10)
        axis.set_xticks([])
        axis.set_yticks([])
        for spine in axis.spines.values():
            spine.set_linewidth(0.8)
            spine.set_color("#222222")
    fig.suptitle("HKL CTF to TSL ANG/H5 conversion validation", fontsize=14)
    fig.savefig(output_path, dpi=220)
    plt.close(fig)


if __name__ == "__main__":
    main()
